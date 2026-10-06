"""Importa a planilha financeira: cria jogadores, mensalidades, caixa e cobranças.

Idempotente: reimportar a mesma planilha atualiza os valores em vez de duplicar.
"""
import zipfile
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import ValidationError
from app.domain.finance_sheet import ParsedSheet, SheetFormatError, clean_player_name, normalize, parse_rows
from app.models.enums import PlayerType, Position
from app.models.finance import CashCategory, CashEntry, CashKind, CollectionItem, FinanceCollection, MonthlyFee
from app.models.player import Player
from app.models.user import User
from app.repositories.settings_repo import SettingsRepository
from app.schemas.finance import ImportResult
from app.services import audit_service

IMPORT_TAG = "(planilha)"


MAX_UNCOMPRESSED = 50 * 1024 * 1024  # proteção contra "zip bomb"
MAX_ROWS = 2000
MAX_COLS = 60


def _check_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise ValidationError("Arquivo inválido: envie a planilha no formato .xlsx") from exc
    if len(infos) > 500 or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
        raise ValidationError("Planilha grande demais para importar")


def read_workbook(content: bytes) -> ParsedSheet:
    # openpyxl usa defusedxml automaticamente quando instalado (bloqueia XML bomb/XXE)
    _check_archive(content)
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # arquivo corrompido / não é xlsx
        raise ValidationError("Arquivo inválido: envie a planilha no formato .xlsx") from exc
    try:
        ws = wb.worksheets[0]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True, max_col=MAX_COLS)):
            if i >= MAX_ROWS:
                raise ValidationError(f"A planilha passa de {MAX_ROWS} linhas")
            rows.append(tuple(row))
        return parse_rows(rows)
    except SheetFormatError as exc:
        raise ValidationError(f"Planilha fora do formato esperado: {exc}") from exc
    finally:
        wb.close()


class FinanceImportService:
    def __init__(self, session: AsyncSession):
        self.session = session

    async def import_sheet(self, sheet: ParsedSheet, actor: User | None) -> ImportResult:
        players = list(await self.session.scalars(select(Player)))
        index: dict[str, Player] = {}
        for p in players:
            for key in filter(None, (p.name, p.nickname)):
                index.setdefault(normalize(key), p)

        created, matched = [], []
        fee_cells = 0
        player_by_label: dict[str, Player] = {}
        for label, cells in sheet.fees.items():
            name, position = clean_player_name(label)
            player = index.get(normalize(name)) or index.get(normalize(label))
            if player is None:
                player = Player(
                    name=name,
                    type=PlayerType.MENSALISTA,
                    primary_position=Position(position) if position else None,
                    active=True,
                )
                self.session.add(player)
                await self.session.flush()
                index[normalize(name)] = player
                created.append(name)
            else:
                matched.append(player.display_name)
            player_by_label[normalize(label)] = player

            for month, cell in cells.items():
                fee = await self.session.scalar(
                    select(MonthlyFee).where(MonthlyFee.player_id == player.id, MonthlyFee.month == month)
                )
                if fee is None:
                    fee = MonthlyFee(player_id=player.id, month=month)
                    self.session.add(fee)
                fee.amount, fee.marker = cell.amount, cell.marker
                fee_cells += 1

        # Lançamentos de caixa vindos da planilha são substituídos a cada importação
        await self.session.execute(delete(CashEntry).where(CashEntry.description.like(f"%{IMPORT_TAG}")))
        entries = [
            CashEntry(month=m, kind=CashKind.ENTRADA, category=CashCategory.DIARISTAS_COLETE,
                      description=f"Diaristas e coletes {IMPORT_TAG}", amount=v)
            for m, v in sheet.extra_income.items()
        ]
        if sheet.current_month:
            if sheet.field_cost:
                entries.append(CashEntry(month=sheet.current_month, kind=CashKind.SAIDA, category=CashCategory.CAMPO,
                                         description=f"Valor do campo {IMPORT_TAG}", amount=sheet.field_cost))
            if sheet.misc_cost:
                entries.append(CashEntry(month=sheet.current_month, kind=CashKind.SAIDA,
                                         category=CashCategory.DIVERSOS,
                                         description=f"Diversos, churrasco, bola, colete {IMPORT_TAG}",
                                         amount=sheet.misc_cost))
        self.session.add_all(entries)

        # "Saldo mês anterior" vira o saldo de abertura a partir do mês corrente da planilha
        settings = await SettingsRepository(self.session).get_current()
        if sheet.opening_balance is not None and sheet.current_month:
            settings.finance_opening_balance = sheet.opening_balance
            settings.finance_opening_month = sheet.current_month

        for parsed in sheet.collections:
            col = await self.session.scalar(select(FinanceCollection).where(FinanceCollection.title == parsed.title))
            if col is None:
                col = FinanceCollection(title=parsed.title, amount_per_person=parsed.amount_per_person)
                self.session.add(col)
                await self.session.flush()
            else:
                col.amount_per_person = parsed.amount_per_person
                await self.session.execute(delete(CollectionItem).where(CollectionItem.collection_id == col.id))
            for item in parsed.items:
                player = player_by_label.get(normalize(item.name)) or index.get(normalize(item.name))
                self.session.add(
                    CollectionItem(
                        collection_id=col.id,
                        player_id=player.id if player else None,
                        name=player.display_name if player else item.name.title(),
                        amount=item.amount if item.amount is not None else parsed.amount_per_person,
                        paid=item.paid,
                    )
                )

        result = ImportResult(
            players_created=created,
            players_matched=matched,
            fee_cells=fee_cells,
            cash_entries=len(entries),
            collections=len(sheet.collections),
            opening_balance=sheet.opening_balance,
            opening_month=sheet.current_month,
        )
        await audit_service.record(
            self.session, user_id=actor.id if actor else None, action="IMPORT", entity="finance",
            after=result.model_dump(mode="json"),
        )
        await self.session.commit()
        return result
