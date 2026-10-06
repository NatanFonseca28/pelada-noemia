"""Exporta as tabelas do banco para .xlsx (uma aba por tabela)."""
import json
from datetime import UTC, date, datetime
from decimal import Decimal
from enum import Enum
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import Table, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import ValidationError
from app.db.base import Base
import app.models  # noqa: F401  (registra todas as tabelas no metadata)
from app.models.user import User
from app.services import audit_service

TZ = ZoneInfo("America/Sao_Paulo")

# Nunca exportar segredos
EXCLUDED_TABLES = {"refresh_tokens"}
EXCLUDED_COLUMNS = {"password_hash", "token_hash"}

TABLE_LABELS = {
    "users": "Usuários",
    "players": "Jogadores",
    "settings": "Configurações",
    "audit_logs": "Auditoria",
    "monthly_fees": "Mensalidades",
    "cash_entries": "Caixa",
    "finance_collections": "Cobranças avulsas",
    "collection_items": "Itens de cobrança",
    "rounds": "Rodadas",
    "attendances": "Presenças",
    "draws": "Sorteios",
    "teams": "Times",
    "team_players": "Jogadores nos times",
    "tournaments": "Campeonatos",
    "tournament_groups": "Grupos",
    "tournament_group_teams": "Times nos grupos",
    "matches": "Partidas",
    "match_events": "Súmula (eventos)",
}

HEADER_FILL = PatternFill("solid", fgColor="027A48")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def exportable_tables() -> list[Table]:
    return [t for t in Base.metadata.sorted_tables if t.name not in EXCLUDED_TABLES]


def _cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, datetime):
        # Excel não aceita fuso; converte para o horário de Brasília
        return value.astimezone(TZ).replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, (date, Decimal, int, float, bool, str)):
        return value
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


class ExportService:
    def __init__(self, session: AsyncSession):
        self.session = session

    async def _name_lookup(self) -> dict[str, dict[int, str]]:
        from app.models.player import Player

        players = {p.id: p.display_name for p in await self.session.scalars(select(Player))}
        users = {u.id: u.name for u in await self.session.scalars(select(User))}
        return {"players": players, "users": users}

    async def export_xlsx(self, table_names: list[str] | None, actor: User) -> bytes:
        tables = exportable_tables()
        if table_names:
            unknown = set(table_names) - {t.name for t in tables}
            if unknown:
                raise ValidationError(f"Tabelas desconhecidas: {', '.join(sorted(unknown))}")
            tables = [t for t in tables if t.name in table_names]

        wb = Workbook()
        info = wb.active
        info.title = "Sobre"
        info.append(["Pelada de Quarta — exportação de dados"])
        info["A1"].font = Font(bold=True, size=14)
        info.append(["Exportado em", datetime.now(TZ).replace(tzinfo=None)])
        info.append(["Exportado por", actor.name])
        info.append([])
        info.append(["Aba", "Tabela", "Registros"])
        for c in info[5]:
            c.font, c.fill = HEADER_FONT, HEADER_FILL
        info.column_dimensions["A"].width = 24
        info.column_dimensions["B"].width = 24
        info.column_dimensions["C"].width = 22

        names = await self._name_lookup()
        for table in tables:
            columns = [c for c in table.columns if c.name not in EXCLUDED_COLUMNS]
            stmt = select(*columns)
            if "id" in table.c:
                stmt = stmt.order_by(table.c.id)
            rows = (await self.session.execute(stmt)).all()

            # Ao lado de cada FK para jogador/usuário, uma coluna com o nome (legível no Excel)
            layout: list[tuple[str, int, str | None]] = []  # (cabeçalho, índice na linha, tabela referenciada)
            for i, col in enumerate(columns):
                layout.append((col.name, i, None))
                ref = next((fk.column.table.name for fk in col.foreign_keys), None)
                if ref in names:
                    layout.append((f"{col.name.removesuffix('_id').removesuffix('_by')}_nome", i, ref))

            title = TABLE_LABELS.get(table.name, table.name)[:31]
            ws = wb.create_sheet(title)
            ws.append([h for h, _, _ in layout])
            for c in ws[1]:
                c.font, c.fill = HEADER_FONT, HEADER_FILL
            for row in rows:
                ws.append([names[ref].get(row[i]) if ref else _cell(row[i]) for _, i, ref in layout])

            for idx, (_, i, ref) in enumerate(layout, start=1):
                col = columns[i]
                letter = get_column_letter(idx)
                values = [str(ws.cell(r, idx).value or "") for r in range(1, min(ws.max_row, 200) + 1)]
                ws.column_dimensions[letter].width = min(max(len(v) for v in values) + 2, 50)
                fmt = None if ref else _number_format(col)
                if fmt:
                    for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                        cell.number_format = fmt
            ws.freeze_panes = "A2"
            if rows:
                ws.auto_filter.ref = ws.dimensions
            info.append([title, table.name, len(rows)])

        await audit_service.record(
            self.session, user_id=actor.id, action="EXPORT", entity="export",
            after={"tables": [t.name for t in tables], "at": datetime.now(UTC).isoformat()},
        )
        await self.session.commit()

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()


def _number_format(column) -> str | None:
    type_name = type(column.type).__name__
    if type_name == "DateTime":
        return "dd/mm/yyyy hh:mm"
    if type_name == "Date":
        return "mm/yyyy" if column.name.endswith("month") else "dd/mm/yyyy"
    if type_name == "Numeric":
        return '"R$" #,##0.00' if column.table.name in {"monthly_fees", "cash_entries", "finance_collections",
                                                        "collection_items", "settings"} else "0.00"
    return None
