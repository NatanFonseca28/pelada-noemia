import csv
import hashlib
import io
import json
import zipfile
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.services.export_service import EXCLUDED_COLUMNS, csv_value, exportable_tables

pytestmark = pytest.mark.asyncio(loop_scope="session")


async def test_exportacao_so_admin(client, mesario_headers, jogador_headers):
    for h in (mesario_headers, jogador_headers):
        assert (await client.get("/api/export/xlsx", headers=h)).status_code == 403
        assert (await client.get("/api/export/csv", headers=h)).status_code == 403
        assert (await client.get("/api/export/tables", headers=h)).status_code == 403


async def _seed(client, h):
    pid = (await client.post("/api/players", json={"name": "=HYPERLINK(\"x\")", "type": "MENSALISTA",
                                                   "primary_position": "ALA", "phone": "21987654321"},
                             headers=h)).json()["id"]
    await client.put("/api/finance/fees", json={"player_id": pid, "month": "2026-10-01", "amount": "37.55"}, headers=h)
    r = await client.post("/api/finance/entries", json={"month": "2026-10-01", "kind": "SAIDA", "category": "CAMPO",
                                                 "description": "Aluguel", "amount": "900.10"}, headers=h)
    assert r.status_code == 201, r.text
    return pid


def _read_zip(content: bytes):
    zf = zipfile.ZipFile(BytesIO(content))
    manifest = json.loads(zf.read("manifest.json"))
    files = {n: zf.read(n) for n in zf.namelist()}
    return manifest, files


async def test_csv_e_copia_fiel_do_banco(client, admin_headers, superadmin_headers, session_factory):
    await _seed(client, admin_headers)
    r = await client.get("/api/export/csv", headers=superadmin_headers)
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    manifest, files = _read_zip(r.content)

    tables = {t.name: t for t in exportable_tables()}
    assert {t["tabela"] for t in manifest["tabelas"]} == set(tables)  # superadmin: todas, inclusive auditoria
    assert "refresh_tokens.csv" not in files

    async with session_factory() as s:
        for entry in manifest["tabelas"]:
            table = tables[entry["tabela"]]
            raw = files[entry["arquivo"]]
            assert hashlib.sha256(raw).hexdigest() == entry["sha256"]
            rows = list(csv.reader(io.StringIO(raw.decode())))
            cols = [c for c in table.columns if c.name not in EXCLUDED_COLUMNS]
            assert rows[0] == [c.name for c in cols] == entry["colunas"]  # colunas reais, na ordem do banco
            db_rows = (await s.execute(select(*cols).order_by(*(list(table.primary_key.columns) or cols)))).all()
            if table.name == "audit_logs":  # o registro da própria exportação é gravado depois da fotografia
                assert db_rows[-1].action == "EXPORT" and db_rows[-1].entity == "export"
                db_rows = db_rows[:-1]
            assert len(rows) - 1 == len(db_rows) == entry["registros"]
            for got, expected in zip(rows[1:], db_rows, strict=True):
                assert got == [csv_value(v) or "" for v in expected], table.name
    sums = files["SHA256SUMS"].decode()
    assert all(f"{t['sha256']}  {t['arquivo']}" in sums for t in manifest["tabelas"])

    fees = list(csv.DictReader(io.StringIO(files["monthly_fees.csv"].decode())))
    assert fees[0]["amount"] == "37.55" and fees[0]["month"] == "2026-10-01"  # decimal exato, data ISO
    players = files["players.csv"].decode()
    assert "+5521987654321" in players and "password_hash" not in files["users.csv"].decode()


async def test_null_diferente_de_texto_vazio_no_csv():
    buf = io.StringIO()
    csv.writer(buf, quoting=csv.QUOTE_NOTNULL).writerow([csv_value(None), csv_value(""), csv_value(Decimal("1.50")),
                                                         csv_value(True), csv_value(datetime(2026, 1, 2, 3, 4, tzinfo=UTC))])
    assert buf.getvalue().strip() == ',"","1.50","true","2026-01-02T03:04:00+00:00"'


async def test_xlsx_fiel_e_com_os_mesmos_hashes_do_csv(client, admin_headers, session_factory):
    await _seed(client, admin_headers)
    x = await client.get("/api/export/xlsx", headers=admin_headers)
    c = await client.get("/api/export/csv", headers=admin_headers)
    wb = load_workbook(BytesIO(x.content))
    manifest, _ = _read_zip(c.content)

    # admin comum: sem auditoria nem log de acessos
    names = {t["tabela"] for t in manifest["tabelas"]}
    assert "audit_logs" not in names and "access_logs" not in names
    assert wb.sheetnames[0] == "_sobre" and set(wb.sheetnames[1:]) == names  # aba = nome real da tabela

    about = {row[0]: row for row in wb["_sobre"].iter_rows(values_only=True) if row and row[0] in names}
    for t in manifest["tabelas"]:
        # mesma fotografia de dados → mesmo hash nos dois formatos (nenhuma escrita entre as duas exportações)
        if t["tabela"] != "monthly_fees":
            continue
        assert about[t["tabela"]][2] == t["registros"] and about[t["tabela"]][4] == t["sha256"]

    fees = wb["monthly_fees"]
    header = [cell.value for cell in fees[1]]
    assert header == [col["colunas"] for col in manifest["tabelas"] if col["tabela"] == "monthly_fees"][0]
    row = dict(zip(header, [cell.value for cell in fees[2]], strict=True))
    assert Decimal(str(row["amount"])) == Decimal("37.55")
    # texto que parece fórmula fica como texto
    players = wb["players"]
    pheader = [cell.value for cell in players[1]]
    name_cell = players.cell(2, pheader.index("name") + 1)
    assert name_cell.value == '=HYPERLINK("x")' and name_cell.data_type == "s"
    assert "player_nome" not in header  # nenhuma coluna inventada


async def test_exporta_tabelas_escolhidas_e_restricoes(client, admin_headers, superadmin_headers):
    r = await client.get("/api/export/xlsx", params=[("tables", "players"), ("tables", "cash_entries")],
                         headers=admin_headers)
    assert load_workbook(BytesIO(r.content)).sheetnames == ["_sobre", "cash_entries", "players"]
    assert (await client.get("/api/export/xlsx", params={"tables": "refresh_tokens"},
                             headers=admin_headers)).status_code == 422
    for fmt in ("xlsx", "csv"):
        r = await client.get(f"/api/export/{fmt}", params={"tables": "audit_logs"}, headers=admin_headers)
        assert r.status_code == 403
    assert (await client.get("/api/export/csv", params={"tables": "audit_logs"},
                             headers=superadmin_headers)).status_code == 200
    listed = {t["name"] for t in (await client.get("/api/export/tables", headers=admin_headers)).json()}
    assert "audit_logs" not in listed and "players" in listed
    listed_super = {t["name"] for t in (await client.get("/api/export/tables", headers=superadmin_headers)).json()}
    assert {"audit_logs", "access_logs"} <= listed_super
    # a exportação fica na auditoria com contagem e hash de cada tabela
    logs = (await client.get("/api/audit", params={"entity": "export"}, headers=superadmin_headers)).json()
    assert any("sha256" in (log["after"]["tables"].get("players") or {}) for log in logs)


async def test_exporta_campeonato_com_sumula(client, admin_headers, mesario_headers, superadmin_headers):
    """Tipos de todas as tabelas (UUID da súmula, JSON, horários) precisam sair nos dois formatos."""
    from tests.api.test_stats import ev, roster, setup_tournament

    _, t = await setup_tournament(client, admin_headers)
    m = t["matches"][0]
    _, rosters = await roster(client, admin_headers, m["id"])
    home = m["home"]["id"]
    r = await client.post(f"/api/matches/{m['id']}/events", json=ev("GOL", home, rosters[home][0]["player_id"]),
                          headers=mesario_headers)
    assert r.status_code in (200, 201), r.text

    x = await client.get("/api/export/xlsx", headers=superadmin_headers)
    assert x.status_code == 200, x.text
    wb = load_workbook(BytesIO(x.content))
    events = wb["match_events"]
    header = [c.value for c in events[1]]
    cid = events.cell(2, header.index("client_event_id") + 1).value
    assert isinstance(cid, str) and len(cid) == 36
    c = await client.get("/api/export/csv", headers=superadmin_headers)
    manifest, files = _read_zip(c.content)
    assert cid in files["match_events.csv"].decode()
    counts = {t["tabela"]: t["registros"] for t in manifest["tabelas"]}
    assert counts["match_events"] == 1 and counts["matches"] >= 1 and counts["team_players"] >= 12
